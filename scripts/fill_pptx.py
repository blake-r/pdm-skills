# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0"]
# ///
"""Fill a tool's template.pptx with data and write a new file.

Usage:
    uv run fill_pptx.py TEMPLATE.pptx DATA.json OUTPUT.pptx [--force] [--list]

Placeholders in the template look like {{key}} or {{key: hint for the author}}.
The whole token is replaced with the value of `key` from DATA.json.

DATA.json:
    {
      "text":   {"key": "value", "other_key": ["line 1", "line 2"]},
      "charts": {"chart_shape_name": {
          "categories": ["Task A", "Task B"],
          "series": [{"name": "Start, day", "values": [0, 10]},
                     {"name": "Duration, days", "values": [10, 5]}],
          "value_axis": {"minimum": 0, "maximum": 21,
                         "major_unit": 7, "number_format": "0"}
      },
      "xy_chart_shape_name": {
          "points": [{"label": "Idea A", "x": 2, "y": 8}]
      }}
    }

For an XY (scatter) chart, "points" replaces every series that shows a marker
with one series per point, named by its label; series without markers
(guide lines, quadrant labels) are kept as they are, unless "lines" gives new
points for them by series name: {"lines": {"Border": [[50, 0], [50, 100]]}}.
"x_axis" sets the horizontal axis of an XY chart with the same keys as
"value_axis".

If the chart's alt text contains "invert_x", its log x axis is mirrored in the
data (plotted as 1/x, clamped to 0.1-10) so that every viewer shows high values
on the left; pass x as is.

A point may carry "size" (e.g. revenue): the marker area is then proportional
to it, the largest point getting the biggest marker.

A list value becomes one paragraph per item. A group shape whose name starts
with "optional" is removed when none of its placeholders received a value.
`--list` prints the placeholder keys and hints of the template and exits.
"""

import argparse
import json
import re
import sys
import textwrap
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_MARKER_STYLE
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

TOKEN = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*(?::([^}]*))?\}\}")
MIN_FONT_PT = 9
MIN_MARKER_PT, MAX_MARKER_PT = 8, 60


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def text_frames(shape):
    if shape.has_text_frame:
        yield shape, shape.text_frame
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield shape, cell.text_frame


def tokens_in(shape):
    found = []
    for s in iter_shapes([shape]) if shape.shape_type == MSO_SHAPE_TYPE.GROUP else [shape]:
        for _, tf in text_frames(s):
            found += [(m.group(1), (m.group(2) or "").strip()) for m in TOKEN.finditer(tf.text)]
    return found


def has_value(values, key):
    v = values.get(key)
    if v is None:
        return False
    if isinstance(v, list):
        return any(str(x).strip() for x in v)
    return bool(str(v).strip())


def as_lines(value):
    if value is None:
        return [""]
    if isinstance(value, list):
        return [str(x) for x in value]
    return str(value).split("\n")


def fill_frame(tf, values, missing):
    text = tf.text
    if not TOKEN.search(text):
        return False
    for key in {m.group(1) for m in TOKEN.finditer(text)}:
        if key not in values:
            missing.add(key)
    whole = TOKEN.fullmatch(text.strip())
    first_par = tf.paragraphs[0]
    first_run = first_par.runs[0] if first_par.runs else None
    font = first_run.font if first_run else None
    style = {
        "size": font.size if font else None,
        "bold": font.bold if font else None,
        "color": font.color.rgb if font and font.color.type == MSO_COLOR_TYPE.RGB else None,
        "name": font.name if font else None,
        "align": first_par.alignment,
        "space_after": first_par.space_after,
    }
    if whole:
        lines = as_lines(values.get(whole.group(1), ""))
    else:
        lines = [TOKEN.sub(lambda m: " ".join(as_lines(values.get(m.group(1), ""))), text)]
    tf.clear()
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = style["align"]
        if style["space_after"] is not None:
            par.space_after = style["space_after"]
        run = par.add_run()
        run.text = line
        if style["size"]:
            run.font.size = style["size"]
        if style["bold"] is not None:
            run.font.bold = style["bold"]
        if style["name"]:
            run.font.name = style["name"]
        if style["color"] is not None:
            run.font.color.rgb = style["color"]
    return True


def shrink_to_fit(shape, tf):
    """Rough fit: lower the font size until the text should fit the box."""
    if not getattr(shape, "width", None) or not shape.height:
        return
    runs = [r for p in tf.paragraphs for r in p.runs]
    if not runs or not runs[0].font.size or not "".join(r.text for r in runs).strip():
        return
    size = runs[0].font.size.pt
    # Average glyph width as a share of the font size; bold Cyrillic runs wide.
    glyph = 0.62 if runs[0].font.bold else 0.56
    width_pt = shape.width / 12700 - 14
    height_pt = shape.height / 12700 - 7.2
    if tf._txBody.bodyPr.get("vert", "").startswith("vert"):  # rotated text runs along the height
        width_pt, height_pt = shape.height / 12700 - 14, shape.width / 12700 - 7.2
    fits = False
    # The floor never goes above the template's own size: a label set at 8 pt is checked at 8 pt.
    floor = min(MIN_FONT_PT, size)
    while size >= floor:
        chars_per_line = max(1, int(width_pt / (size * glyph)))
        words_too_long = any(len(w) > chars_per_line for p in tf.paragraphs for w in p.text.split())
        lines = sum(
            len(textwrap.wrap(p.text, chars_per_line, break_long_words=True)) or 1
            for p in tf.paragraphs
        )
        if not words_too_long and lines * size * 1.2 <= height_pt:
            fits = True
            break
        size -= 0.5
    if not fits:
        size = floor
        preview = (shape.name + ": " + "".join(r.text for p in tf.paragraphs for r in p.runs))[:80]
        print(f"warning: text does not fit the box even at {floor} pt, shorten it: {preview!r}", file=sys.stderr)
    for r in runs:
        r.font.size = Pt(size)


def fill_xy_chart(chart, points, lines=None):
    lines = lines or {}
    data = XyChartData()
    for ser in chart.plots[0].series:
        if ser.marker.style != XL_MARKER_STYLE.NONE:
            continue
        kept = data.add_series(ser.name)
        if ser.name in lines:
            for x, y in lines[ser.name]:
                kept.add_data_point(x, y)
            continue
        xs = ser._element.xpath("./c:xVal//c:pt/c:v")
        ys = ser._element.xpath("./c:yVal//c:pt/c:v")
        for x, y in zip(xs, ys):
            kept.add_data_point(float(x.text), float(y.text))
    for point in points:
        data.add_series(point["label"]).add_data_point(point["x"], point["y"])
    chart.replace_data(data)
    sizes = {p["label"]: p["size"] for p in points if p.get("size") is not None}
    if sizes:
        top = max(sizes.values())
        for ser in chart.plots[0].series:
            if ser.name in sizes and top > 0:
                ser.marker.size = max(MIN_MARKER_PT, round(MAX_MARKER_PT * (sizes[ser.name] / top) ** 0.5))


def invert_x(spec):
    """Mirror a log x axis for charts whose alt text says "invert_x".

    Viewers disagree on reversed axes, so such a template plots 1/x on a fixed
    0.1-10 log axis with its own tick labels; data keep their natural x.
    """
    if "x_axis" in spec:
        sys.exit("this chart has a fixed mirrored x axis: drop 'x_axis'")

    def flip(x):
        return 1 / min(max(float(x), 0.1), 10)

    spec = dict(spec)
    spec["points"] = [{**p, "x": flip(p["x"])} for p in spec["points"]]
    spec["lines"] = {k: [[flip(x), y] for x, y in v] for k, v in spec.get("lines", {}).items()}
    return spec


def fill_chart(shape, spec):
    chart = shape.chart
    if "points" in spec and "invert_x" in (shape._element.xpath("./*[1]/p:cNvPr/@descr") or [""])[0]:
        spec = invert_x(spec)
    if "points" in spec:
        unknown = set(spec.get("lines", {})) - {ser.name for ser in chart.plots[0].series}
        if unknown:
            sys.exit(f"no series named: {', '.join(sorted(unknown))}")
        fill_xy_chart(chart, spec["points"], spec.get("lines"))
    else:
        data = CategoryChartData()
        data.categories = spec["categories"]
        for series in spec["series"]:
            data.add_series(series["name"], series["values"])
        chart.replace_data(data)
    for key, axis in (("value_axis", chart.value_axis), ("x_axis", chart.category_axis)):
        axis_spec = spec.get(key) or {}
        for attr in ("minimum", "maximum", "major_unit"):
            if attr in axis_spec:
                setattr(axis, f"{attr}_scale" if attr != "major_unit" else attr, axis_spec[attr])
        if "number_format" in axis_spec:
            axis.tick_labels.number_format = axis_spec["number_format"]
            axis.tick_labels.number_format_is_linked = False


def list_placeholders(prs):
    seen = {}
    for slide in prs.slides:
        for shape in iter_shapes(slide.shapes):
            if shape.has_chart:
                seen.setdefault(f"chart:{shape.name}", "native chart, fill via 'charts'")
            for _, tf in text_frames(shape):
                for m in TOKEN.finditer(tf.text):
                    seen.setdefault(m.group(1), (m.group(2) or "").strip())
    for key, hint in seen.items():
        print(f"{key}\t{hint}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("template")
    ap.add_argument("data", nargs="?")
    ap.add_argument("output", nargs="?")
    ap.add_argument("--force", action="store_true", help="overwrite OUTPUT")
    ap.add_argument("--list", action="store_true", help="print placeholders and exit")
    args = ap.parse_args()

    prs = Presentation(args.template)
    if args.list:
        list_placeholders(prs)
        return
    if not args.data or not args.output:
        ap.error("DATA and OUTPUT are required unless --list is given")
    out = Path(args.output)
    if out.exists() and not args.force:
        sys.exit(f"{out} exists; pass --force to overwrite")

    data = json.loads(Path(args.data).read_text(encoding="utf-8"))
    values = data.get("text", {})
    charts = data.get("charts", {})
    missing = set()
    sample_charts = []

    for slide in prs.slides:
        for shape in list(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP and shape.name.startswith("optional"):
                keys = [k for k, _ in tokens_in(shape)]
                if keys and not any(has_value(values, k) for k in keys):
                    shape._element.getparent().remove(shape._element)
        for shape in iter_shapes(slide.shapes):
            if shape.has_chart:
                if shape.name in charts:
                    fill_chart(shape, charts.pop(shape.name))
                else:
                    sample_charts.append(shape.name)
            for owner, tf in text_frames(shape):
                if fill_frame(tf, values, missing) and owner is shape and shape.has_text_frame:
                    shrink_to_fit(shape, tf)

    if charts:
        sys.exit(f"no chart shapes named: {', '.join(charts)}")
    prs.save(out)
    if missing:
        print(f"warning: no value for {', '.join(sorted(missing))}; left empty", file=sys.stderr)
    if sample_charts:
        print(f"warning: chart {', '.join(sample_charts)} keeps the template's sample data; add it under 'charts'", file=sys.stderr)
    print(out)


if __name__ == "__main__":
    main()
