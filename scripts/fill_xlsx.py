# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Fill a tool's template.xlsx with data and write a new file.

Usage:
    uv run fill_xlsx.py TEMPLATE.xlsx DATA.json OUTPUT.xlsx [--force] [--list]

The template marks its data table with the defined name `data_start`
(the first cell of the first data row). Header cells carry hints as comments.
Formula columns are prepared in the template; for rows beyond the prepared
ones the formulas of the first data row are copied down.

DATA.json:
    {
      "text": {"title": "Q3 prioritization"},   # replaces {{key}} tokens in cells; a list becomes lines
      "cells": {"B1": "value", "Sheet!B2": 5},   # optional direct writes; no sheet means the data sheet
      "rows": [["Feature A", 1200, 2, 0.8, 3]]   # written from data_start
    }

`--list` prints the header row with its hints and exits.
"""

import argparse
from copy import copy
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

TOKEN = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*(?::([^}]*))?\}\}")


def data_start(wb):
    name = wb.defined_names.get("data_start")
    if name is None:
        return wb.active, 1, 2
    sheet_name, ref = next(iter(name.destinations))
    col, row = coordinate_from_string(ref.replace("$", ""))
    return wb[sheet_name], column_index_from_string(col), row


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("template")
    ap.add_argument("data", nargs="?")
    ap.add_argument("output", nargs="?")
    ap.add_argument("--force", action="store_true", help="overwrite OUTPUT")
    ap.add_argument("--list", action="store_true", help="print header hints and exit")
    args = ap.parse_args()

    wb = load_workbook(args.template)
    ws, col0, row0 = data_start(wb)

    if args.list:
        for cell in ws[row0 - 1]:
            if cell.value is not None:
                hint = cell.comment.text.strip() if cell.comment else ""
                print(f"{cell.column_letter}\t{cell.value}\t{hint}")
        return
    if not args.data or not args.output:
        ap.error("DATA and OUTPUT are required unless --list is given")
    out = Path(args.output)
    if out.exists() and not args.force:
        sys.exit(f"{out} exists; pass --force to overwrite")

    data = json.loads(Path(args.data).read_text(encoding="utf-8"))
    values = data.get("text", {})
    missing = set()

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and TOKEN.search(cell.value):
                    def sub(m):
                        if m.group(1) not in values:
                            missing.add(m.group(1))
                        v = values.get(m.group(1), "")
                        return "\n".join(map(str, v)) if isinstance(v, list) else str(v)
                    cell.value = TOKEN.sub(sub, cell.value)
                    if "\n" in cell.value:
                        wrapped = copy(cell.alignment)
                        wrapped.wrap_text = True
                        cell.alignment = wrapped

    for ref, value in data.get("cells", {}).items():
        sheet_name, _, cell_ref = ref.rpartition("!")
        (wb[sheet_name.strip("'")] if sheet_name else ws)[cell_ref] = value

    formulas = {
        c.column: c.value
        for c in ws[row0]
        if isinstance(c.value, str) and c.value.startswith("=")
    }
    for i, values_row in enumerate(data.get("rows", [])):
        r = row0 + i
        for j, value in enumerate(values_row):
            col = col0 + j
            if col in formulas:
                sys.exit(f"row {i + 1}: column {ws.cell(r, col).column_letter} holds a formula; drop that value")
            ws.cell(r, col, value)
        for col, formula in formulas.items():
            target = ws.cell(r, col)
            if target.value is None:
                origin = ws.cell(row0, col).coordinate
                target.value = Translator(formula, origin=origin).translate_formula(target.coordinate)
                target.number_format = ws.cell(row0, col).number_format

    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    if missing:
        print(f"warning: no value for {', '.join(sorted(missing))}; left empty", file=sys.stderr)
    print(out)


if __name__ == "__main__":
    main()
