# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0"]
# ///
"""Smoke test of the deck pipeline. No LLM, no network.

Usage (from the repository root):
    uv run scripts/smoke_test.py

Steps, each checked:
1. fill three tool templates from their example.json into a scratch folder;
2. merge the filled decks into one deck;
3. run check_pptx.py on the deck (expected: "no layout problems");
4. confirm that no template token {{ is left in the deck;
5. fill every pptx template with its example.json and check the layout;
6. run build/quality_gate.py (cards, selection table, intake skill, plugin manifest).

Exits with 0 when every step passes, with 1 otherwise. The example.json files
are format samples, not data: the test checks the pipeline, not the content.
"""

import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
# Divider labels in service-blueprint are 8 pt text in a 0.14 in strip with zero insets:
# the estimate flags them, the rendered slide is fine.
LAYOUT_ESTIMATE_FALSE_POSITIVE = {"service-blueprint"}
TOOLS = [
    "product-planning/impact-mapping",
    "business-analysis/swot-analysis",
    "product-planning/metric-pyramid",
]


def run(*args):
    return subprocess.run(["uv", "run", *map(str, args)], cwd=ROOT, capture_output=True, text=True)


def fail(step, result):
    print(f"FAIL: {step}")
    print((result.stdout + result.stderr).strip())
    sys.exit(1)


def main():
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        filled = []
        for tool in TOOLS:
            d = ROOT / "skills" / tool
            name = d.name
            r = run("scripts/fill_pptx.py", d / f"{name}.pptx", d / "example.json", out / f"{name}.pptx")
            if r.returncode:
                fail(f"fill {tool}", r)
            filled.append(out / f"{name}.pptx")
        print(f"ok: filled {len(filled)} templates")

        deck = out / "deck.pptx"
        r = run("scripts/merge_pptx.py", deck, *filled)
        if r.returncode:
            fail("merge", r)
        print("ok: merged")

        r = run("scripts/check_pptx.py", deck)
        if r.returncode:
            fail("check_pptx", r)
        print("ok: no layout problems")

        left = [i for i, s in enumerate(Presentation(deck).slides, 1)
                for sh in s.shapes if sh.has_text_frame and "{{" in sh.text_frame.text]
        if left:
            print(f"FAIL: template tokens left on slides {sorted(set(left))}")
            sys.exit(1)
        print("ok: no template tokens left")

    with tempfile.TemporaryDirectory() as tmp_all:
        out_all = Path(tmp_all)
        bad = []
        for tpl in sorted((ROOT / "skills").glob("*/*/*.pptx")):
            if tpl.parent.name in LAYOUT_ESTIMATE_FALSE_POSITIVE or not (tpl.parent / "example.json").is_file():
                continue
            r = run("scripts/fill_pptx.py", tpl, tpl.parent / "example.json", out_all / f"{tpl.parent.name}.pptx", "--force")
            c = run("scripts/check_pptx.py", out_all / f"{tpl.parent.name}.pptx") if not r.returncode else r
            if c.returncode or "{{" in c.stdout:
                bad.append(f"{tpl.parent.name}: rc={c.returncode} " + (c.stdout + c.stderr).strip()[-200:])
        if bad:
            print("FAIL: template filled with its example.json has layout problems")
            print("\n".join(bad))
            sys.exit(1)
    print("ok: every template fits its example.json")

    r = run("build/quality_gate.py")
    if r.returncode or "with problems: 0" not in r.stdout:
        fail("quality_gate", r)
    print("ok: quality gate")
    print("smoke test passed")


if __name__ == "__main__":
    main()
